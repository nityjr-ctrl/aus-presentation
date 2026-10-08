#!/usr/bin/env python3
"""Structural checks for AUS-teaching-deck-3d.pptx (and test-one-slide.pptx). No PowerPoint needed.

    python validate.py            # writes validation-3d.json beside this script; exit 1 on any failure
"""
import json
import posixpath
import re
import struct
import sys
import zipfile
from pathlib import Path

from lxml import etree

HERE = Path(__file__).resolve().parent
PRES = HERE.parent
ORIG = PRES / "AUS-teaching-deck.pptx"
DECK = PRES / "AUS-teaching-deck-3d.pptx"
TEST = HERE / "test-one-slide.pptx"
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "am3d": "http://schemas.microsoft.com/office/drawing/2017/model3d",
    "p159": "http://schemas.microsoft.com/office/powerpoint/2015/09/main",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}
R_EMBED = f"{{{NS['r']}}}embed"
DISCLAIMER = "Teaching illustration. Not patient-specific. Not a medical device."
TERMS = ["DICOM", "patient", "MRN", "NHS number"]
MAX_BYTES = 3 * 1024 * 1024
MAX_TRIS = 100_000

results = []


def check(name, ok, detail=""):
    results.append({"check": name, "pass": bool(ok), "detail": detail})
    return ok


class Pkg:
    def __init__(self, path):
        self.z = zipfile.ZipFile(path)
        self.names = set(self.z.namelist())
        self.xml = {}

    def read(self, name):
        return self.z.read(name)

    def tree(self, name):
        if name not in self.xml:
            self.xml[name] = etree.fromstring(self.read(name))
        return self.xml[name]

    def rels(self, part):
        d, f = posixpath.split(part)
        rp = posixpath.join(d, "_rels", f + ".rels")
        out = {}
        if rp not in self.names:
            return out
        for r in self.tree(rp):
            if r.get("TargetMode") == "External":
                continue
            t = r.get("Target")
            tgt = t.lstrip("/") if t.startswith("/") else posixpath.normpath(posixpath.join(d, t))
            out[r.get("Id")] = (r.get("Type"), tgt)
        return out

    def slides(self):
        pres = self.tree("ppt/presentation.xml")
        rels = self.rels("ppt/presentation.xml")
        return [rels[s.get(f"{{{NS['r']}}}id")][1] for s in pres.find("p:sldIdLst", NS)]

    def notes_of(self, slide):
        for t, tgt in self.rels(slide).values():
            if t.endswith("/notesSlide"):
                return tgt
        return None


def texts(root, skip_names=()):
    out = []
    for sp in root.iter(f"{{{NS['p']}}}sp", f"{{{NS['p']}}}graphicFrame", f"{{{NS['p']}}}pic"):
        c = sp.find(".//p:cNvPr", NS)
        if c is not None and c.get("name") in skip_names:
            continue
        if sp.getparent() is not None and etree.QName(sp.getparent()).localname in ("sp",):
            continue
        out.extend(t.text or "" for t in sp.iter(f"{{{NS['a']}}}t"))
    return out


def glb_info(data):
    magic, _, length = struct.unpack_from("<III", data, 0)
    clen, ctype = struct.unpack_from("<II", data, 12)
    js = json.loads(data[20:20 + clen])
    tris = 0
    for m in js.get("meshes", []):
        for p in m["primitives"]:
            acc = js["accessors"][p["indices"]] if "indices" in p else js["accessors"][p["attributes"]["POSITION"]]
            tris += acc["count"] // 3
    nodes = [n.get("name") for n in js.get("nodes", [])]
    return {"bytes": len(data), "triangles": tris, "nodes": nodes, "animations": len(js.get("animations", [])),
            "copyright": js.get("asset", {}).get("copyright", ""), "json": data[20:20 + clen].decode("utf-8")}


def validate_package(path, label, steps, full):
    pkg = Pkg(path)
    # 1. every XML part parses
    bad = []
    xml_parts = [n for n in pkg.names if n.endswith(".xml") or n.endswith(".rels")]
    for n in xml_parts:
        try:
            pkg.tree(n)
        except etree.XMLSyntaxError as e:
            bad.append(f"{n}: {e}")
    check(f"{label}: all {len(xml_parts)} XML and .rels parts parse", not bad, "; ".join(bad))
    # 2. every relationship resolves
    missing = []
    nrels = 0
    for n in pkg.names:
        if not n.endswith(".rels"):
            continue
        src = n.replace("_rels/", "")[:-5]
        for rid, (t, tgt) in pkg.rels(src).items():
            nrels += 1
            if tgt not in pkg.names:
                missing.append(f"{n} {rid} -> {tgt}")
    check(f"{label}: all {nrels} internal relationships resolve", not missing, "; ".join(missing[:10]))
    # 3. content types
    ct = pkg.tree("[Content_Types].xml")
    defaults = {d.get("Extension").lower(): d.get("ContentType") for d in ct.findall("ct:Default", NS)}
    overrides = {o.get("PartName").lstrip("/"): o.get("ContentType") for o in ct.findall("ct:Override", NS)}
    need = {"glb": "model/gltf.binary", "png": "image/png"}
    if full:
        need["mp4"] = "video/mp4"
    check(f"{label}: content-type defaults for {', '.join(need)}", all(defaults.get(k) == v for k, v in need.items()),
          json.dumps({k: defaults.get(k) for k in need}))
    sl = [n for n in pkg.names if re.fullmatch(r"ppt/(slides|notesSlides)/\w+\.xml", n)]
    no_ct = [n for n in sl if n not in overrides]
    check(f"{label}: overrides for all {len(sl)} slide and notes parts", not no_ct, ", ".join(no_ct))
    check(f"{label}: [Content_Types].xml is the first zip entry", pkg.z.namelist()[0] == "[Content_Types].xml")
    # 4. the 3D frames
    order = pkg.slides()
    frames = 0
    glbs = {}
    for st in steps:
        pos = st["slide"] if full else 1
        slide = order[pos - 1]
        root = pkg.tree(slide)
        rels = pkg.rels(slide)
        models = root.findall(".//am3d:model3d", NS)
        ok = len(models) == 1
        detail = f"slide {st['slide']} ({slide}): {len(models)} model3d"
        if ok:
            m = models[0]
            t, glb = rels.get(m.get(R_EMBED), (None, None))
            blip = m.find("am3d:raster/am3d:blip", NS)
            _, png = rels.get(blip.get(R_EMBED) if blip is not None else "", (None, None))
            ac = m.xpath("ancestor::mc:AlternateContent", namespaces=NS)[0]
            fb = ac.find("mc:Fallback/p:pic", NS)
            fb_png = rels.get(fb.find(".//a:blip", NS).get(R_EMBED), (None, None))[1] if fb is not None else None
            names = [c.get("name") for c in ac.iter(f"{{{NS['p']}}}cNvPr")]
            ok = (t and t.endswith("/model3d") and glb and glb.endswith(".glb") and glb in pkg.names and png
                  and png.endswith(".png") and png in pkg.names and fb_png == png and names == ["!!AUS3D", "!!AUS3D"]
                  and ac.find("mc:Choice", NS).get("Requires") == "am3d")
            order_ok = [etree.QName(c).localname for c in m] [:4] == ["spPr", "camera", "trans", "raster"]
            ok = ok and order_ok
            detail += f", glb {glb}, raster {png}, fallback {fb_png}, names {names}, child order ok {order_ok}"
            if glb in pkg.names:
                glbs[st["id"]] = glb_info(pkg.read(glb))
            disc = DISCLAIMER in texts(root)
            ok = ok and disc
            detail += f", disclaimer {disc}"
            if full:
                tr = root.find("mc:AlternateContent", NS)
                has_morph = tr is not None and tr.find(".//p159:morph", NS) is not None and tr.find("mc:Fallback/p:transition/p:fade", NS) is not None
                ok = ok and (has_morph == st["morph"])
                detail += f", morph {has_morph} (expected {st['morph']})"
        frames += int(bool(ok))
        check(f"{label}: 3D frame on slide {st['slide']}", ok, detail)
    # 5. GLB budgets and attribution
    for sid, g in glbs.items():
        check(f"{label}: {sid}.glb under 3 MB and 100k triangles", g["bytes"] < MAX_BYTES and g["triangles"] < MAX_TRIS,
              f"{g['bytes']} bytes, {g['triangles']} triangles, {len(g['nodes'])} nodes, {g['animations']} animations")
        check(f"{label}: {sid}.glb carries the attribution", "CC-BY-SA 4.0" in g["copyright"] and "Not patient-specific" in g["copyright"])
    return pkg, order, frames, glbs


def patient_data_scan(pkg, orig):
    """New occurrences of the terms are allowed only inside the added disclaimer, notes and GLB attribution."""
    allowed = [DISCLAIMER, "Not patient-specific", "notPatientSpecific"]
    def count(z, strip_allowed):
        totals = {t: 0 for t in TERMS}
        for n in z.namelist():
            data = z.read(n)
            if n.lower().endswith((".png", ".jpg", ".jpeg", ".mp4")):
                if len(data) > 132 and data[128:132] == b"DICM":
                    totals["DICOM"] += 1000
                continue
            s = data.decode("utf-8", "ignore")
            if strip_allowed:
                # the 3D frames' alt text repeats the existing slide title plus the disclaimer
                s = re.sub(r'descr="3D model: [^"]*"', "", s)
                for a in allowed:
                    s = s.replace(a, "")
            for t in TERMS:
                totals[t] += len(re.findall(re.escape(t), s, re.IGNORECASE if t != "MRN" else 0))
        return totals
    new = count(pkg.z, True)
    old = count(zipfile.ZipFile(orig), False)
    ok = all(new[t] <= old[t] for t in TERMS) and new["DICOM"] == 0 and new["MRN"] == 0 and new["NHS number"] == 0
    check("deck: no patient data (no DICOM files or MRN/NHS number strings; 'patient' occurs no more often than in the original deck, outside the added disclaimers and alt text)",
          ok, f"new {new}, original {old}")


def text_integrity(pkg, order, orig, steps):
    o = Pkg(orig)
    oorder = o.slides()
    added_notes = {s["slide"]: s["note"] for s in steps}
    bad = []
    for i, oslide in enumerate(oorder, start=1):
        ni = i if i <= 39 else i + 1
        nslide = order[ni - 1]
        a = texts(o.tree(oslide))
        b = texts(pkg.tree(nslide), skip_names=("AUS3D disclaimer",))
        if a != b:
            bad.append(f"slide {i}: text runs differ")
        on, nn = o.notes_of(oslide), pkg.notes_of(nslide)
        at = [t.text or "" for t in o.tree(on).iter(f"{{{NS['a']}}}t")] if on else []
        bt = [t.text or "" for t in pkg.tree(nn).iter(f"{{{NS['a']}}}t")] if nn else []
        expect = at + ([added_notes[i]] if i in added_notes else [])
        if bt != expect:
            bad.append(f"slide {i}: notes differ")
    check(f"deck: text runs of all {len(oorder)} original slides and their notes unchanged (only the disclaimer and one notes line added)",
          not bad, "; ".join(bad))


def video_slide(pkg, order):
    slide = order[39]
    root = pkg.tree(slide)
    t = texts(root)
    rels = pkg.rels(slide)
    vids = [tgt for typ, tgt in rels.values() if typ.endswith("/video") or typ.endswith("/media")]
    has_vf = root.find(".//a:videoFile", NS) is not None
    ok = (t[:1] == ["Backup: UroOps 3D AUS flow (video)"] and "Backup: outside the timed talk" in t and has_vf
          and vids and all(v in pkg.names and v.endswith(".mp4") for v in vids) and len(order) == 50)
    check("deck: backup video slide at position 40 with a resolvable MP4 and the backup footer", ok,
          f"{slide}: texts {t}, media {sorted(set(vids))}, slides {len(order)}")


def python_pptx_load(path, expect):
    from pptx import Presentation
    try:
        n = len(Presentation(str(path)).slides)
        check(f"python-pptx opens {path.name} ({n} slides)", n == expect, f"expected {expect}")
    except Exception as e:  # noqa: BLE001
        check(f"python-pptx opens {path.name}", False, repr(e))


def main():
    steps = json.loads((HERE / "steps.json").read_text(encoding="utf-8"))["slides"]
    for st in steps:
        f = HERE / "steps" / f"{st['id']}.glb"
        g = glb_info(f.read_bytes())
        check(f"steps/{f.name}: under budget", g["bytes"] < MAX_BYTES and g["triangles"] < MAX_TRIS, f"{g['bytes']} bytes, {g['triangles']} triangles")
    pkg, order, frames, glbs = validate_package(DECK, "deck", steps, True)
    check("deck: seven 3D frames", frames == 7, f"{frames}")
    patient_data_scan(pkg, ORIG)
    text_integrity(pkg, order, ORIG, steps)
    video_slide(pkg, order)
    test_step = [s for s in steps if s["slide"] == 28]
    tpkg, torder, tframes, _ = validate_package(TEST, "test-one-slide", test_step, False)
    check("test-one-slide: exactly one slide with one 3D frame", len(torder) == 1 and tframes == 1, f"{len(torder)} slides")
    python_pptx_load(DECK, 50)
    python_pptx_load(TEST, 1)
    passed = sum(r["pass"] for r in results)
    report = {"passed": passed, "failed": len(results) - passed, "checks": results,
              "glbs": {k: {kk: vv for kk, vv in v.items() if kk != "json"} for k, v in glbs.items()}}
    (HERE / "validation-3d.json").write_text(json.dumps(report, indent=1) + "\n", encoding="utf-8")
    for r in results:
        print(("PASS " if r["pass"] else "FAIL ") + r["check"] + ("" if r["pass"] else f"  [{r['detail']}]"))
    print(f"{passed} passed, {len(results) - passed} failed")
    sys.exit(0 if passed == len(results) else 1)


if __name__ == "__main__":
    main()
