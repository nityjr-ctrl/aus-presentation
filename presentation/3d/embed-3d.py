#!/usr/bin/env python3
"""Build presentation/AUS-teaching-deck-3d.pptx from the unchanged AUS-teaching-deck.pptx.

For every slide in steps.json: remove the static UroOps screenshot (where flagged), insert a native
PowerPoint 3D model (am3d, see AM3D-NOTES.md) with its PNG fallback, add the disclaimer text box,
append one speaker-note paragraph and, where flagged, a Morph transition with a fade fallback.
Then add one backup video slide after slide 39, and write test-one-slide.pptx (slide 28 only).

Rebuild from committed models, PNGs and the backup video in the existing 3D deck:
    python embed-3d.py

Use the optional Node scripts only when deliberately regenerating visual assets.
"""
import argparse
import hashlib
import json
import posixpath
import re
import shutil
import tempfile
import uuid
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape, quoteattr

from lxml import etree

HERE = Path(__file__).resolve().parent
PRES = HERE.parent
SRC = PRES / "AUS-teaching-deck.pptx"
OUT = PRES / "AUS-teaching-deck-3d.pptx"
TEST = HERE / "test-one-slide.pptx"
VIDEO = HERE / "video" / "aus-flow-1080p.mp4"
POSTER = HERE / "video" / "aus-flow-poster.png"
VIDEO_SHA256 = "aea54cc59ed74e87e72b0194e76cb95f49329a0277419853ba6997bb2d9c2982"

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "am3d": "http://schemas.microsoft.com/office/drawing/2017/model3d",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}
REL_MODEL3D = "http://schemas.microsoft.com/office/2017/06/relationships/model3d"
REL_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
REL_NOTES = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide"
EMU_IN = 914400
EMU_PER_M = 36_000_000
DISCLAIMER = "Teaching illustration. Not patient-specific. Not a medical device."
VIDEO_TITLE = "Backup: UroOps 3D AUS flow (video)"
VIDEO_NOTE = ("Backup only, outside the timed talk: a silent screen recording of the public UroOps 3D AUS lesson "
              "(uroops3d.com/lab/aus), all eleven steps and then the final configuration. Use it if the 3D models "
              "do not load. Teaching schematic, clinical review pending.")


def emu(inches):
    return int(round(inches * EMU_IN))


def parse(path):
    return etree.parse(str(path), etree.XMLParser(remove_blank_text=False))


def write(tree, path):
    tree.write(str(path), xml_declaration=True, encoding="UTF-8", standalone=True)


def rels_path(part):
    return part.parent / "_rels" / f"{part.name}.rels"


def add_rel(rels_file, rtype, target):
    tree = parse(rels_file)
    root = tree.getroot()
    ids = {r.get("Id") for r in root}
    n = 1
    while f"rIdAus3d{n}" in ids:
        n += 1
    rid = f"rIdAus3d{n}"
    etree.SubElement(root, f"{{{NS['pr']}}}Relationship", Id=rid, Type=rtype, Target=target)
    write(tree, rels_file)
    return rid


def ensure_default(ct_file, ext, ctype):
    tree = parse(ct_file)
    root = tree.getroot()
    if not root.xpath(f"ct:Default[@Extension='{ext}']", namespaces=NS):
        el = etree.Element(f"{{{NS['ct']}}}Default", Extension=ext, ContentType=ctype)
        root.insert(0, el)
    write(tree, ct_file)


def frame_xml(shape_id, name, descr, box, rid_glb, rid_png, cam, trans, animated):
    x, y, cx, cy = emu(box["left"]), emu(box["top"]), emu(box["width"]), emu(box["height"])
    guid = "{" + str(uuid.uuid5(uuid.NAMESPACE_URL, f"aus3d/{descr}")).upper() + "}"
    pos, look = cam["pos"], cam["lookAt"]
    anim = ""
    if animated:
        anim = (
            '<am3d:extLst>'
            '<a:ext uri="{9A65AA19-BECB-4387-8358-8AD5134E1D82}">'
            '<a3danim:embedAnim xmlns:a3danim="http://schemas.microsoft.com/office/drawing/2018/animation/model3d" animId="0">'
            '<a3danim:animPr length="8000" count="indefinite"/></a3danim:embedAnim></a:ext>'
            '<a:ext uri="{E9DE012E-A134-456F-84FE-255F9AAD75C6}">'
            '<a3danim:posterFrame xmlns:a3danim="http://schemas.microsoft.com/office/drawing/2018/animation/model3d" animId="0"/>'
            '</a:ext></am3d:extLst>'
        )
    lights = "".join(
        f'<am3d:ptLight rad="0"><am3d:clr><a:scrgbClr r="{r}" g="{g}" b="{b}"/></am3d:clr>'
        f'<am3d:intensity n="{i}" d="1000000"/><am3d:pos x="{px}" y="{py}" z="{pz}"/></am3d:ptLight>'
        for r, g, b, i, px, py, pz in [
            (100000, 75000, 50000, 9765625, 21959998, 70920001, 16344003),
            (40000, 60000, 95000, 12250000, -37964106, 51130435, 57631972),
            (86837, 72700, 100000, 3125000, -37739122, 58056624, -34769649),
        ]
    )
    cnvpr = (
        f'<p:cNvPr id="{shape_id}" name={quoteattr(name)} descr={quoteattr(descr)}>'
        '<a:extLst><a:ext uri="{FF2B5EF4-FFF2-40B4-BE49-F238E27FC236}">'
        f'<a16:creationId xmlns:a16="http://schemas.microsoft.com/office/drawing/2014/main" id="{guid}"/>'
        '</a:ext></a:extLst></p:cNvPr>'
    )
    xml = (
        f'<mc:AlternateContent xmlns:mc="{NS["mc"]}" xmlns:p="{NS["p"]}" xmlns:a="{NS["a"]}" xmlns:r="{NS["r"]}">'
        f'<mc:Choice xmlns:am3d="{NS["am3d"]}" Requires="am3d">'
        '<p:graphicFrame>'
        f'<p:nvGraphicFramePr>{cnvpr}<p:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></p:cNvGraphicFramePr><p:nvPr/></p:nvGraphicFramePr>'
        f'<p:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></p:xfrm>'
        f'<a:graphic><a:graphicData uri="{NS["am3d"]}">'
        f'<am3d:model3d r:embed="{rid_glb}">'
        f'<am3d:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></am3d:spPr>'
        '<am3d:camera>'
        f'<am3d:pos x="{pos[0]}" y="{pos[1]}" z="{pos[2]}"/>'
        f'<am3d:up dx="0" dy="{EMU_PER_M}" dz="0"/>'
        f'<am3d:lookAt x="{look[0]}" y="{look[1]}" z="{look[2]}"/>'
        f'<am3d:perspective fov="{cam["fov"]}"/>'
        '</am3d:camera>'
        '<am3d:trans>'
        f'<am3d:meterPerModelUnit n="{trans["n"]}" d="1000000"/>'
        f'<am3d:preTrans dx="{trans["pre"][0]}" dy="{trans["pre"][1]}" dz="{trans["pre"][2]}"/>'
        '<am3d:scale><am3d:sx n="1000000" d="1000000"/><am3d:sy n="1000000" d="1000000"/><am3d:sz n="1000000" d="1000000"/></am3d:scale>'
        '<am3d:rot/>'
        '<am3d:postTrans dx="0" dy="0" dz="0"/>'
        '</am3d:trans>'
        f'<am3d:raster rName="Office3DRenderer" rVer="16.0.8326"><am3d:blip r:embed="{rid_png}"/></am3d:raster>'
        f'{anim}'
        '<am3d:winViewport/>'
        '<am3d:ambientLight><am3d:clr><a:scrgbClr r="50000" g="50000" b="50000"/></am3d:clr><am3d:illuminance n="500000" d="1000000"/></am3d:ambientLight>'
        f'{lights}'
        '</am3d:model3d></a:graphicData></a:graphic>'
        '</p:graphicFrame>'
        '</mc:Choice>'
        '<mc:Fallback>'
        '<p:pic>'
        f'<p:nvPicPr>{cnvpr}<p:cNvPicPr><a:picLocks noGrp="1" noRot="1" noChangeAspect="1" noMove="1" noResize="1" noEditPoints="1" noAdjustHandles="1" noChangeArrowheads="1" noChangeShapeType="1" noCrop="1"/></p:cNvPicPr><p:nvPr/></p:nvPicPr>'
        f'<p:blipFill><a:blip r:embed="{rid_png}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
        f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
        '</p:pic>'
        '</mc:Fallback>'
        '</mc:AlternateContent>'
    )
    return etree.fromstring(xml)


def textbox_xml(shape_id, name, box, text, size_hundredths, colour="526773"):
    x, y, cx, cy = emu(box["left"]), emu(box["top"]), emu(box["width"]), emu(box["height"])
    return etree.fromstring(
        f'<p:sp xmlns:p="{NS["p"]}" xmlns:a="{NS["a"]}">'
        f'<p:nvSpPr><p:cNvPr id="{shape_id}" name="{name}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></p:spPr>'
        '<p:txBody><a:bodyPr spcFirstLastPara="1" wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t"><a:noAutofit/></a:bodyPr><a:lstStyle/>'
        '<a:p><a:pPr marL="0" marR="0" lvl="0" indent="0" algn="l"><a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="0"/></a:spcAft><a:buNone/></a:pPr>'
        f'<a:r><a:rPr lang="en-GB" sz="{size_hundredths}" b="0"><a:solidFill><a:srgbClr val="{colour}"/></a:solidFill><a:latin typeface="Arial"/><a:ea typeface="Arial"/><a:cs typeface="Arial"/></a:rPr>'
        f'<a:t>{escape(text)}</a:t></a:r><a:endParaRPr/></a:p></p:txBody></p:sp>'
    )


def note_paragraph(text):
    return etree.fromstring(
        f'<a:p xmlns:a="{NS["a"]}"><a:pPr marL="0" lvl="0" indent="0" algn="l"><a:spcBef><a:spcPts val="0"/></a:spcBef>'
        '<a:spcAft><a:spcPts val="0"/></a:spcAft><a:buNone/></a:pPr>'
        f'<a:r><a:rPr lang="en-GB" sz="1200"/><a:t>{escape(text)}</a:t></a:r><a:endParaRPr/></a:p>'
    )


def morph_xml():
    return etree.fromstring(
        f'<mc:AlternateContent xmlns:mc="{NS["mc"]}" xmlns:p="{NS["p"]}">'
        '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
        '<p:transition xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" spd="slow" p14:dur="1500">'
        '<p159:morph option="byObject"/></p:transition>'
        '</mc:Choice>'
        '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback>'
        '</mc:AlternateContent>'
    )


def max_shape_id(root):
    ids = [int(v) for v in root.xpath("//p:cNvPr/@id", namespaces=NS) if v.isdigit()]
    return max(ids) if ids else 1


def embed_slide(d, step, k, frame):
    slide = d / "ppt" / "slides" / f"slide{step['slide']}.xml"
    tree = parse(slide)
    root = tree.getroot()
    sptree = root.find("p:cSld/p:spTree", NS)
    srels = rels_path(slide)
    # 1. the static UroOps screenshot, and its relationship
    insert_at = len(sptree)
    if step["removePicture"]:
        pics = sptree.findall("p:pic", NS)
        if len(pics) != 1:
            raise SystemExit(f"slide {step['slide']}: expected one picture, found {len(pics)}")
        pic = pics[0]
        rid = pic.find(".//a:blip", NS).get(f"{{{NS['r']}}}embed")
        insert_at = list(sptree).index(pic)
        sptree.remove(pic)
        if rid not in etree.tostring(root).decode():
            rt = parse(srels)
            for rel in rt.getroot():
                if rel.get("Id") == rid:
                    rt.getroot().remove(rel)
            write(rt, srels)
    # 2. media parts and relationships (absolute targets, as the rest of this package uses)
    media = d / "ppt" / "media"
    glb_name = f"model3d{k}.glb"
    png_name = f"aus-fallback-{step['slide']}.png"
    shutil.copyfile(HERE / "steps" / f"{step['id']}.glb", media / glb_name)
    shutil.copyfile(HERE / "renders" / f"{step['id']}.png", media / png_name)
    rid_glb = add_rel(srels, REL_MODEL3D, f"/ppt/media/{glb_name}")
    rid_png = add_rel(srels, REL_IMAGE, f"/ppt/media/{png_name}")
    # 3. the 3D frame, in the picture's place in the z-order
    sid = max_shape_id(root) + 1
    r = step["camera"]["resolved"]
    to_emu = lambda p: [int(round((p[i] - frame["centre"][i]) * frame["mpu"] * EMU_PER_M)) for i in range(3)]
    cam = {"pos": to_emu(r["posMetres"]), "lookAt": to_emu(r["focusMetres"]), "fov": int(r["fovDegrees"] * 60000)}
    descr = f"3D model: {step['title']}. {DISCLAIMER}"
    ac = frame_xml(sid, "!!AUS3D", descr, step["frame"], rid_glb, rid_png, cam, frame["trans"], step["state"]["keepAnimation"])
    sptree.insert(insert_at, ac)
    # 4. the disclaimer text box
    sptree.append(textbox_xml(sid + 1, "AUS3D disclaimer", step["disclaimer"], DISCLAIMER, 1100, step.get("disclaimerColour", "526773")))
    # slide 32 only: narrow the row text boxes so their glyphs clear the frame (text unchanged)
    narrow = step.get("narrowTextBoxes")
    if narrow:
        for shp in sptree.findall("p:sp", NS):
            if int(shp.find(".//p:cNvPr", NS).get("id")) in narrow["shapeIds"]:
                shp.find("p:spPr/a:xfrm/a:ext", NS).set("cx", str(emu(narrow["width"])))
    # 5. Morph on the incoming slide, fade fallback (these slides have no clrMapOvr or timing)
    if step["morph"]:
        if root.find("p:transition", NS) is not None or root.find("mc:AlternateContent", NS) is not None:
            raise SystemExit(f"slide {step['slide']} already has a transition")
        csld = root.find("p:cSld", NS)
        csld.addnext(morph_xml())
    write(tree, slide)
    # 6. one speaker-note paragraph
    rt = parse(srels)
    notes_target = [x.get("Target") for x in rt.getroot() if x.get("Type") == REL_NOTES][0]
    notes = d / notes_target.lstrip("/") if notes_target.startswith("/") else (slide.parent / notes_target).resolve()
    nt = parse(notes)
    body = nt.getroot().xpath("//p:sp[.//p:ph[@type='body']]/p:txBody", namespaces=NS)[0]
    body.append(note_paragraph(step["note"]))
    write(nt, notes)
    return {"slide": step["slide"], "shape_id": sid, "glb": glb_name, "png": png_name, "camera": cam}


def target_part(source, target):
    """Resolve either an absolute or relative OPC relationship target."""
    return (target.lstrip("/") if target.startswith("/") else
            posixpath.normpath(posixpath.join(posixpath.dirname(source), target)))


def keep_slides(d, names):
    """Trim both the presentation list and its slide relationships before pruning."""
    pres = d / "ppt" / "presentation.xml"
    tree = parse(pres)
    rt = parse(rels_path(pres))
    keep = set()
    for rel in list(rt.getroot()):
        if rel.get("Type").endswith("/slide"):
            if Path(rel.get("Target")).name in names:
                keep.add(rel.get("Id"))
            else:
                rt.getroot().remove(rel)
    for slide in list(tree.getroot().find("p:sldIdLst", NS)):
        if slide.get(f"{{{NS['r']}}}id") not in keep:
            slide.getparent().remove(slide)
    write(tree, pres)
    write(rt, rels_path(pres))


def prune_package(d):
    """Keep parts reachable from package relationships, including notes and media.

    This replaces the session-local clean.py. External links are left untouched;
    missing internal targets fail the build instead of creating a broken PPTX.
    """
    retained = {"[Content_Types].xml"}
    pending = [""]
    visited = set()
    while pending:
        source = pending.pop()
        if source in visited:
            continue
        visited.add(source)
        if source:
            retained.add(source)
        rp = rels_path(d / source) if source else d / "_rels" / ".rels"
        if not rp.exists():
            continue
        retained.add(rp.relative_to(d).as_posix())
        for rel in parse(rp).getroot():
            if rel.get("TargetMode") == "External":
                continue
            target = target_part(source, rel.get("Target"))
            if target.startswith("../") or not (d / target).is_file():
                raise ValueError(f"missing internal relationship: {source} -> {target}")
            pending.append(target)
    ct = parse(d / "[Content_Types].xml")
    for entry in list(ct.getroot()):
        if entry.tag == f"{{{NS['ct']}}}Override" and entry.get("PartName").lstrip("/") not in retained:
            ct.getroot().remove(entry)
    write(ct, d / "[Content_Types].xml")
    for path in d.rglob("*"):
        if path.is_file() and path.relative_to(d).as_posix() not in retained:
            path.unlink()


def duplicate_backup_template(d):
    """Copy slide 40 after slide 39, with its shared layout, using native OOXML.

    The template's notes are deliberately not copied: add_movie supplies the
    backup narration. Original slides, notes and timing parts are never edited.
    """
    slides = d / "ppt" / "slides"
    number = max(int(p.stem[5:]) for p in slides.glob("slide*.xml")) + 1
    new = f"slide{number}.xml"
    path = slides / new
    shutil.copyfile(slides / "slide40.xml", path)
    rt = parse(rels_path(slides / "slide40.xml"))
    for rel in list(rt.getroot()):
        if not rel.get("Type").endswith("/slideLayout"):
            rt.getroot().remove(rel)
    write(rt, rels_path(path))
    pres = d / "ppt" / "presentation.xml"
    rid = add_rel(rels_path(pres), NS["r"] + "/slide", f"/ppt/slides/{new}")
    tree = parse(pres)
    slide_list = tree.getroot().find("p:sldIdLst", NS)
    sid = max(int(s.get("id")) for s in slide_list) + 1
    entry = etree.Element(f"{{{NS['p']}}}sldId", id=str(sid))
    entry.set(f"{{{NS['r']}}}id", rid)
    slide_list.insert(39, entry)
    write(tree, pres)
    ct_path = d / "[Content_Types].xml"
    ct = parse(ct_path)
    template = ct.getroot().xpath("ct:Override[@PartName='/ppt/slides/slide40.xml']", namespaces=NS)[0]
    etree.SubElement(ct.getroot(), f"{{{NS['ct']}}}Override",
                     PartName=f"/ppt/slides/{new}", ContentType=template.get("ContentType"))
    write(ct, ct_path)
    return new


def resolve_video(video, seed, tmp):
    """Reuse the exact versioned MP4 if no standalone recording is available."""
    if video.is_file():
        return video
    if not seed.is_file():
        raise ValueError(f"missing video {video}; supply --video MP4 or --video-from an existing 3D deck")
    with zipfile.ZipFile(seed) as z:
        matches = [n for n in z.namelist() if n.lower().endswith(".mp4")]
        if len(matches) != 1:
            raise ValueError(f"expected one embedded MP4 in {seed}, found {len(matches)}; supply --video")
        data = z.read(matches[0])
    if hashlib.sha256(data).hexdigest() != VIDEO_SHA256:
        raise ValueError(f"embedded backup video checksum mismatch in {seed}; supply --video for a new recording")
    result = tmp / "aus-flow-1080p.mp4"
    result.write_bytes(data)
    return result


def zip_dir(d, out):
    tmp = out.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(d / "[Content_Types].xml", "[Content_Types].xml")
        for f in sorted(d.rglob("*")):
            if f.is_file() and f.name != "[Content_Types].xml":
                z.write(f, f.relative_to(d).as_posix())
    tmp.replace(out)


def slide_order(d):
    pres = (d / "ppt" / "presentation.xml").read_text(encoding="utf-8-sig")
    rels = (d / "ppt" / "_rels" / "presentation.xml.rels").read_text(encoding="utf-8-sig")
    rid = {m.group(2): m.group(1) for m in re.finditer(r'Target="([^"]+)"[^>]*Id="([^"]+)"', rels)}
    rid.update({m.group(1): m.group(2) for m in re.finditer(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels)})
    return [rid[i].split("/")[-1] for i in re.findall(r'<p:sldId [^>]*r:id="([^"]+)"', pres)]


def add_video_slide(d):
    order = slide_order(d)
    assert order[38] == "slide39.xml" and order[39] == "slide40.xml", order[38:40]
    new = duplicate_backup_template(d)
    path = d / "ppt" / "slides" / new
    tree = parse(path)
    root = tree.getroot()
    sptree = root.find("p:cSld/p:spTree", NS)
    for gf in sptree.findall("p:graphicFrame", NS):
        sptree.remove(gf)
    texts = ["".join(sp.xpath(".//a:t/text()", namespaces=NS)) for sp in sptree.findall("p:sp", NS)]
    title = sptree.findall("p:sp", NS)[texts.index("Outcome evidence: device and endpoint")]
    runs = title.findall(".//a:r", NS)
    runs[0].find("a:t", NS).text = VIDEO_TITLE
    for extra in runs[1:]:
        extra.getparent().remove(extra)
    assert "Backup: outside the timed talk" in texts, texts
    write(tree, path)
    # drop the duplicated table's relationships, if any
    rp = rels_path(path)
    rt = parse(rp)
    for rel in list(rt.getroot()):
        if not rel.get("Type").endswith("/slideLayout"):
            rt.getroot().remove(rel)
    write(rt, rp)
    return new


def add_movie(pptx, video):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation(str(pptx))
    slide = prs.slides[39]
    title = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text == VIDEO_TITLE]
    assert title, "video slide not at position 40"
    slide.shapes.add_movie(str(video), Inches(0.67), Inches(1.9), Inches(8.9), Inches(5.0),
                           poster_frame_image=str(POSTER), mime_type="video/mp4")
    slide.notes_slide.notes_text_frame.text = VIDEO_NOTE
    prs.save(str(pptx))


def finalise_content_types(pptx):
    """python-pptx writes overrides for glb and mp4; add the Defaults too and keep [Content_Types].xml first."""
    with tempfile.TemporaryDirectory() as t:
        t = Path(t)
        with zipfile.ZipFile(pptx) as z:
            z.extractall(t)
        ct = t / "[Content_Types].xml"
        ensure_default(ct, "glb", "model/gltf.binary")
        ensure_default(ct, "mp4", "video/mp4")
        ensure_default(ct, "png", "image/png")
        zip_dir(t, pptx)


def scene_frame(cfg):
    sf = cfg["sceneFrame"]
    half = max((sf["max"][i] - sf["min"][i]) / 2 for i in range(3))
    n = int(round(1e6 / (2 * half)))
    mpu = n / 1e6
    pre = [int(round(-c * mpu * EMU_PER_M)) for c in sf["centre"]]
    return {"centre": sf["centre"], "mpu": mpu, "trans": {"n": n, "pre": pre}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUT)
    parser.add_argument("--test-output", type=Path, default=TEST)
    parser.add_argument("--video", type=Path, default=VIDEO, help="Standalone MP4, if available")
    parser.add_argument("--video-from", type=Path, default=OUT,
                        help="Existing 3D deck containing the versioned backup MP4")
    args = parser.parse_args()
    if args.output.resolve() == SRC.resolve() or args.test_output.resolve() == SRC.resolve():
        parser.error("output paths must not overwrite the original teaching deck")
    if args.output.resolve() == args.test_output.resolve():
        parser.error("output and test-output must be different files")
    cfg = json.loads((HERE / "steps.json").read_text(encoding="utf-8"))
    frame = scene_frame(cfg)
    for need in [SRC, POSTER, *[HERE / "steps" / f"{s['id']}.glb" for s in cfg["slides"]],
                 *[HERE / "renders" / f"{s['id']}.png" for s in cfg["slides"]]]:
        if not need.exists():
            raise SystemExit(f"missing {need}")
    with tempfile.TemporaryDirectory() as tmp:
        try:
            video = resolve_video(args.video, args.video_from, Path(tmp))
        except ValueError as exc:
            parser.error(str(exc))
        d = Path(tmp) / "deck"
        with zipfile.ZipFile(SRC) as z:
            z.extractall(d)
        assert slide_order(d) == [f"slide{i}.xml" for i in range(1, 50)], "unexpected slide order"
        ensure_default(d / "[Content_Types].xml", "glb", "model/gltf.binary")
        done = [embed_slide(d, step, k, frame) for k, step in enumerate(cfg["slides"], start=1)]
        # one-slide smoke test: slide 28 only
        t = Path(tmp) / "test"
        shutil.copytree(d, t)
        keep_slides(t, {"slide28.xml"})
        prune_package(t)
        zip_dir(t, args.test_output)
        # the backup video slide after slide 39, then tidy and pack
        new = add_video_slide(d)
        prune_package(d)
        # Stage the full deck until video insertion succeeds, leaving any seed intact.
        built = Path(tmp) / "built.pptx"
        zip_dir(d, built)
        add_movie(built, video)
        finalise_content_types(built)
        shutil.copyfile(built, args.output)
    print(json.dumps({"embedded": done, "videoSlide": new, "deck": str(args.output), "test": str(args.test_output)}, indent=1))


if __name__ == "__main__":
    main()
